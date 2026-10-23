"""References retain original provenance; free-form documents are never auto-approved specs."""

import json
import math
import re
from pathlib import Path
from .documents import extract_document


def extract_reference(path: Path, file_type: str) -> dict:
    fields, conflicts, specification, requirements = [], [], {}, []
    if file_type in ("application/json", "json"):
        data = json.loads(
            Path(path).read_text(),
            parse_constant=lambda value: (_ for _ in ()).throw(
                ValueError("Non-finite JSON number")
            ),
        )
        if len(Path(path).read_bytes()) > 1_000_000:
            raise ValueError("Reference JSON too large")
        if not isinstance(data, dict) or data.get("schemaVersion") != "1":
            raise ValueError("Reference requires schemaVersion 1")
        specification = data.get("product", {})
        requirements = data.get("requiredStatements", [])
        if not isinstance(specification, dict) or not isinstance(requirements, list):
            raise ValueError("Invalid product or requiredStatements")
        quantity = specification.get("netQuantity")
        if quantity is not None and (
            not isinstance(quantity, dict)
            or isinstance(quantity.get("value"), bool)
            or not isinstance(quantity.get("value"), (int, float))
            or not math.isfinite(quantity["value"])
            or quantity["value"] <= 0
            or not isinstance(quantity.get("unit"), str)
        ):
            raise ValueError("Invalid netQuantity")
        if quantity and quantity["unit"].lower() not in (
            "ml",
            "l",
            "g",
            "kg",
            "oz",
            "fl oz",
        ):
            raise ValueError("Unsupported quantity unit")
        for req in requirements:
            if not isinstance(req, dict) or not all(
                isinstance(req.get(k), str) and req[k].strip() for k in ("id", "text")
            ):
                raise ValueError("Required statements need id and text")
            alternatives = req.get("anyOf", [req["text"]])
            if not isinstance(alternatives, list) or not 1 <= len(alternatives) <= 20 or not all(isinstance(v, str) and v.strip() and len(v) <= 2000 for v in alternatives):
                raise ValueError("anyOf requires 1 to 20 nonempty disclosure strings")
            condition = req.get("appliesWhen")
            if condition is not None:
                if not isinstance(condition, dict) or not condition or set(condition) - {"juiceBelowPercent", "measurementBasis", "market", "category"}:
                    raise ValueError("Unsupported requirement condition")
                if any(not isinstance(condition[k], str) or not condition[k].strip() for k in ("market", "category", "measurementBasis") if k in condition):
                    raise ValueError("Invalid policy context")
                if "juiceBelowPercent" in condition:
                    value = condition["juiceBelowPercent"]
                    if isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 <= value <= 100 or condition.get("measurementBasis") != "volume":
                        raise ValueError("Juice threshold requires a volume basis and percent from 0 to 100")
                elif "measurementBasis" in condition:
                    raise ValueError("Measurement basis requires a juice threshold")


        def flatten(value, location="$", depth=0):
            if depth > 20:
                raise ValueError("Reference nesting exceeds limit")
            if isinstance(value, dict):
                for key, item in value.items():
                    flatten(item, location + "." + key, depth + 1)
            elif isinstance(value, list):
                for i, item in enumerate(value):
                    flatten(item, f"{location}[{i}]", depth + 1)
            else:
                fields.append(
                    {
                        "key": location.split(".")[-1],
                        "value": value,
                        "location": location,
                    }
                )

        flatten(data)
        text = json.dumps(data, ensure_ascii=False)
        method, status = "structured-json-v1", "read"
    elif file_type in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ):
        from openpyxl import load_workbook
        from zipfile import ZipFile

        with ZipFile(path) as archive:
            if sum(f.file_size for f in archive.infolist()) > 50_000_000:
                raise ValueError("Workbook decompressed size exceeds limit")
        workbook = load_workbook(path, read_only=True, data_only=False)
        lines, by_key = [], {}
        for sheet in workbook:
            if sheet.max_row * sheet.max_column > 100_000:
                raise ValueError("Workbook cell count exceeds limit")
            for row in sheet:
                nonempty = [c for c in row if c.value is not None]
                for cell in nonempty:
                    fields.append(
                        {
                            "key": str(nonempty[0].value),
                            "value": cell.value,
                            "location": f"'{sheet.title}'!{cell.coordinate}",
                        }
                    )
                    lines.append(f"'{sheet.title}'!{cell.coordinate}: {cell.value}")
                if len(nonempty) >= 2:
                    key = str(nonempty[0].value).strip().casefold()
                    value = str(nonempty[1].value)
                    if key in by_key and by_key[key]["value"] != value:
                        conflicts.append(
                            {
                                "key": key,
                                "values": [
                                    by_key[key],
                                    {
                                        "value": value,
                                        "location": f"'{sheet.title}'!{nonempty[1].coordinate}",
                                    },
                                ],
                            }
                        )
                    by_key[key] = {
                        "value": value,
                        "location": f"'{sheet.title}'!{nonempty[1].coordinate}",
                    }
        quantity_fields = []
        for item in fields:
            if re.fullmatch(
                r"net\s*(?:quantity|weight|content|contents)", item["key"].strip(), re.I
            ):
                m = re.fullmatch(
                    r"\s*(\d+(?:\.\d+)?)\s*(mL|L|g|kg|oz|fl oz)\s*",
                    str(item["value"]),
                    re.I,
                )
                if m:
                    quantity_fields.append(
                        {
                            "value": float(m[1]),
                            "unit": m[2],
                            "location": item["location"],
                        }
                    )
        if len(quantity_fields) == 1:
            specification["netQuantity"] = quantity_fields[0]
        elif len(quantity_fields) > 1:
            conflicts.append({"key": "netQuantity", "values": quantity_fields})
        workbook.close()
        text = "\n".join(lines)
        method, status = "xlsx-cells", "read"
    elif file_type in ("application/pdf", "pdf"):
        doc = extract_document(path, file_type)
        text, method, status = doc["text"], doc["method"], doc["readingStatus"]
        quantities = []
        for i, page in enumerate(doc["pages"]):
            fields.append(
                {"key": "pageText", "value": page["text"], "location": f"Page {i + 1}"}
            )
            for match in re.finditer(
                r"\bnet\s*(?:quantity|weight|contents?)\s*[:：]?\s*(\d+(?:\.\d+)?)\s*(fl\s+oz|mL|kg|oz|g|L)\b",
                page["text"],
                re.I,
            ):
                quantities.append(
                    {
                        "value": float(match[1]),
                        "unit": re.sub(r"\s+", " ", match[2]),
                        "location": f"Page {i + 1}: {match[0]}",
                    }
                )
        if len(quantities) == 1:
            specification["netQuantity"] = quantities[0]
        elif len(quantities) > 1:
            conflicts.append({"key": "netQuantity", "values": quantities})
    else:
        raise ValueError("Unsupported reference type")
    return {
        "text": text,
        "fields": fields,
        "specification": specification,
        "requirements": requirements,
        "conflicts": conflicts,
        "readingStatus": status,
        "method": method,
        "qualifiers": [
            word
            for word in ("draft", "reconstructed", "unapproved")
            if word in text.casefold()
        ],
    }

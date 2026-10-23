import asyncio
import subprocess
import sys
import hashlib
import io
import json
import os
import warnings
import zipfile
from urllib.parse import urlparse
import httpx
from pathlib import Path
from uuid import uuid4
from PIL import Image
from pypdf import PdfReader
from app.errors import APIError

MAX_BYTES = 20 * 1024 * 1024
MAX_PIXELS = 24_000_000
Image.MAX_IMAGE_PIXELS = MAX_PIXELS


def storage_root():
    root = Path(
        os.environ.get(
            "WAVELENGTH_STORAGE_DIR",
            "/tmp/wavelength-files" if os.environ.get("VERCEL") else "./private-files",
        )
    ).resolve()
    if "public" in root.parts:
        raise RuntimeError("File storage must be outside public directories")
    root.mkdir(parents=True, exist_ok=True, mode=0o700)
    return root


def validate(data, kind):
    if not data:
        raise APIError(422, "invalid_file", "The uploaded file is empty")
    if len(data) > MAX_BYTES:
        raise APIError(413, "oversized_input", "Files must be at most 20 MiB")
    pages = []
    try:
        if data.startswith(b"%PDF-"):
            pdf = PdfReader(io.BytesIO(data), strict=True)
            if pdf.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported")
            if not 1 <= len(pdf.pages) <= 20:
                raise ValueError("PDFs must contain 1–20 pages")
            for p in pdf.pages:
                width, height = float(p.cropbox.width), float(p.cropbox.height)
                if not 0 < width <= 14400 or not 0 < height <= 14400:
                    raise ValueError("PDF page dimensions are outside supported bounds")
                pages.append(
                    dict(
                        width=width, height=height, rotation=int(p.rotation or 0) % 360
                    )
                )
            return "application/pdf", pages
        if kind == "artwork":
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(data)) as image:
                    if image.format not in ("JPEG", "PNG"):
                        raise ValueError("Artwork must be PDF, PNG, or JPEG")
                    if (
                        image.width * image.height > MAX_PIXELS
                        or max(image.size) > 20000
                    ):
                        raise ValueError(
                            "Image exceeds 24 megapixels or 20,000 pixels per side"
                        )
                    file_type = Image.MIME[image.format]
                    pages = [dict(width=image.width, height=image.height, rotation=0)]
                    image.verify()
            return file_type, pages
        if data.startswith(b"PK"):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if (
                    len(archive.infolist()) > 2000
                    or sum(i.file_size for i in archive.infolist()) > 60 * 1024 * 1024
                ):
                    raise ValueError("Workbook decompressed size exceeds limits")
                if (
                    "[Content_Types].xml" not in archive.namelist()
                    or "xl/workbook.xml" not in archive.namelist()
                ):
                    raise ValueError("Not a supported XLSX workbook")
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                if len(book.worksheets) > 30 or any(
                    s.max_row > 20000 or s.max_column > 200 for s in book
                ):
                    raise ValueError("Workbook dimensions exceed limits")
            finally:
                book.close()
            return (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                [],
            )
        parsed = json.loads(data.decode("utf-8-sig"))
        if not isinstance(parsed, dict):
            raise ValueError("Reference JSON must contain an object")
        return "application/json", []
    except APIError:
        raise
    except Exception as exc:
        raise APIError(
            422,
            "invalid_file",
            "The file could not be validated",
            {"reason": str(exc)[:300]},
        ) from exc


async def save_upload(upload, kind):
    data = await upload.read(MAX_BYTES + 1)
    path = storage_root() / str(uuid4())
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, "wb") as f:
        f.write(data)
    try:
        process = await asyncio.create_subprocess_exec(
            sys.executable,
            "-m",
            "app.processing",
            "validate",
            cwd=str(Path(__file__).resolve().parents[1]),
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        try:
            stdout, stderr = await asyncio.wait_for(
                process.communicate(
                    json.dumps({"path": str(path), "kind": kind}).encode()
                ),
                timeout=30,
            )
        except BaseException:
            process.kill()
            await process.communicate()
            raise
        try:
            response = json.loads(stdout)
        except Exception:
            raise APIError(
                422, "invalid_file", "File validation failed within resource limits"
            )
        if process.returncode or not response.get("ok"):
            raise APIError(
                422,
                "invalid_file",
                "The file could not be validated",
                {"reason": response.get("error", "")[:300]},
            )
        result = response["result"]
        return (
            path,
            hashlib.sha256(data).hexdigest(),
            result["fileType"],
            result["pages"],
        )
    except BaseException:
        path.unlink(missing_ok=True)
        raise


def cloud_storage():
    return os.environ.get("WAVELENGTH_STORAGE_MODE", "").strip().strip("\"'") == "blob"


def download_blob(url, review_id=None, kind=None):
    parsed = urlparse(url)
    host = os.environ.get("WAVELENGTH_BLOB_HOST", "").strip().strip("\"'")
    if (
        not host
        or parsed.scheme != "https"
        or parsed.netloc != host
        or parsed.username
        or parsed.password
        or parsed.query
        or parsed.fragment
    ):
        raise APIError(422, "invalid_file", "Unrecognized private storage object")
    if review_id and not parsed.path.startswith(f"/reviews/{review_id}/{kind}/"):
        raise APIError(404, "not_found", "Object does not belong to this review")
    token = os.environ.get("BLOB_READ_WRITE_TOKEN", "").strip().strip("\"'")
    if not token:
        raise APIError(503, "storage_unavailable", "Private storage is not configured")
    try:
        with httpx.stream(
            "GET",
            url,
            headers={"Authorization": "Bearer " + token},
            timeout=15,
            follow_redirects=False,
        ) as response:
            if response.status_code != 200:
                raise APIError(
                    503,
                    "storage_unavailable",
                    "Private object could not be read",
                    retryable=True,
                )
            content = bytearray()
            for chunk in response.iter_bytes():
                content.extend(chunk)
                if len(content) > MAX_BYTES:
                    raise APIError(
                        413, "oversized_input", "Files must be at most 20 MiB"
                    )
            return bytes(content)
    except httpx.HTTPError:
        raise APIError(
            503,
            "storage_unavailable",
            "Private storage is temporarily unavailable",
            retryable=True,
        )
